"""Davriy tozalashdan (cleanup.py) OLDIN ma'lumotlarni Excel (.xlsx)
formatida zaxiralab qo'yadigan modul.

Har bir bo'lim (xavfsizlik jurnali, yangiliklar, buyurtmalar) uchun ALOHIDA
.xlsx fayl yaratiladi, shunda admin faqat kerakli bo'limni yuklab olishi
mumkin. Fayl admin panelidan BIR MARTA yuklab olingach, serverdan avtomatik
o'chiriladi -- shu tufayli zaxiralar diskda cheksiz to'planib qolmaydi.

Agar admin bir necha kun davomida hech qaysi backupni yuklab olmasa (masalan
band bo'lgani uchun), disk to'lib ketmasligi uchun har bir bo'lim uchun
faqat oxirgi 3 ta zaxira saqlanadi -- undan eskilari avtomatik o'chiriladi.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional

from openpyxl import Workbook

from . import config, models

MAX_BACKUPS_PER_SECTION = 3

SECTION_LABELS = {
    "havfsizlik_jurnali": "Xavfsizlik jurnali",
    "yangiliklar": "Yangiliklar",
    "buyurtmalar": "Buyurtmalar",
}


def _timestamp() -> str:
    return datetime.utcnow().strftime("%Y%m%d_%H%M%S")


def _section_dir(section: str) -> Path:
    d = config.BACKUPS_DIR / section
    d.mkdir(parents=True, exist_ok=True)
    return d


def _enforce_section_limit(section: str) -> None:
    """Bitta bo'limda MAX_BACKUPS_PER_SECTION dan ortiq fayl to'planib
    qolsa, eng eskilarini o'chiradi (disk to'lib ketmasligi uchun)."""
    directory = _section_dir(section)
    files = sorted(directory.glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    while len(files) >= MAX_BACKUPS_PER_SECTION:
        oldest = files.pop(0)
        try:
            oldest.unlink()
        except OSError:
            pass


def _write_workbook(section: str, sheets: "list[tuple[str, list[str], list[list]]]") -> Optional[Path]:
    """`sheets` = [(varaq_nomi, ustun_sarlavhalari, qator_ro'yxati), ...].
    Hech bo'lmaganda bitta qator bo'lmasa, fayl yaratilmaydi (bo'sh
    zaxiraning keragi yo'q)."""
    if not any(rows for _title, _headers, rows in sheets):
        return None

    _enforce_section_limit(section)

    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title=title[:31])
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for col_idx, header in enumerate(headers, start=1):
            width = max(12, min(50, len(str(header)) + 4))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    filename = f"{section}_{_timestamp()}.xlsx"
    path = _section_dir(section) / filename
    wb.save(path)
    return path


def backup_security_logs(events: Iterable["models.SecurityEvent"]) -> Optional[Path]:
    headers = ["ID", "Turi", "XSS turi", "IP manzil", "Yo'l", "Metod", "Namuna", "User-Agent", "Vaqt (UTC)"]
    rows = [
        [
            e.id, e.kind, e.xss_type or "", e.ip_address, e.path, e.method,
            e.matched_sample, e.user_agent or "", e.created_at.isoformat() if e.created_at else "",
        ]
        for e in events
    ]
    return _write_workbook("havfsizlik_jurnali", [("Xavfsizlik jurnali", headers, rows)])


def backup_news(posts: Iterable["models.NewsPost"]) -> Optional[Path]:
    headers = ["ID", "Sarlavha", "Qisqa tavsif", "Matn", "Rasm URL", "E'lon qilinganmi", "Yaratilgan", "Yangilangan"]
    rows = [
        [
            n.id, n.title, n.excerpt or "", n.body or "", n.image_url or "",
            "Ha" if n.is_published else "Yo'q",
            n.created_at.isoformat() if n.created_at else "",
            n.updated_at.isoformat() if n.updated_at else "",
        ]
        for n in posts
    ]
    return _write_workbook("yangiliklar", [("Yangiliklar", headers, rows)])


def backup_orders(orders: Iterable["models.Order"]) -> Optional[Path]:
    orders = list(orders)
    order_headers = [
        "ID", "Kod", "Mijoz ismi", "Holat", "Izoh",
        "Jami summa", "Jami tannarx", "Jami foyda",
        "Yaratilgan", "Tasdiqlangan",
    ]
    order_rows = [
        [
            o.id, o.code, o.customer_name, o.status, o.admin_note or "",
            o.total_amount, o.total_cost, o.total_profit,
            o.created_at.isoformat() if o.created_at else "",
            o.confirmed_at.isoformat() if o.confirmed_at else "",
        ]
        for o in orders
    ]

    item_headers = ["Buyurtma kodi", "Mahsulot", "Soni", "Narxi (dona)", "Tannarxi (dona)", "Jami"]
    item_rows = []
    for o in orders:
        for it in o.items:
            item_rows.append([o.code, it.product_name, it.quantity, it.unit_price, it.unit_cost, it.line_total])

    return _write_workbook(
        "buyurtmalar",
        [("Buyurtmalar", order_headers, order_rows), ("Mahsulotlar", item_headers, item_rows)],
    )


def list_backups() -> List[dict]:
    """Barcha bo'limlardagi mavjud (hali yuklab olinmagan) zaxira
    fayllarini, eng yangisidan boshlab qaytaradi."""
    out: List[dict] = []
    if not config.BACKUPS_DIR.exists():
        return out
    for section_dir in config.BACKUPS_DIR.iterdir():
        if not section_dir.is_dir():
            continue
        section = section_dir.name
        for f in section_dir.glob("*.xlsx"):
            stat = f.stat()
            out.append(
                {
                    "filename": f.name,
                    "section": section,
                    "section_label": SECTION_LABELS.get(section, section),
                    "created_at": datetime.utcfromtimestamp(stat.st_mtime),
                    "size_kb": round(stat.st_size / 1024, 1),
                }
            )
    out.sort(key=lambda x: x["created_at"], reverse=True)
    return out


def resolve_backup_path(filename: str) -> Optional[Path]:
    """Berilgan fayl nomini, faqat BACKUPS_DIR ichidan xavfsiz qidiradi
    (yo'l-traversal urinishlaridan himoyalangan: faqat oddiy fayl nomi
    qabul qilinadi, papka belgilari yo'q)."""
    safe_name = Path(filename).name
    if safe_name != filename or not safe_name.endswith(".xlsx"):
        return None
    for section_dir in config.BACKUPS_DIR.iterdir():
        if not section_dir.is_dir():
            continue
        candidate = section_dir / safe_name
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def delete_backup_file(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except OSError:
        pass
